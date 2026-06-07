# Copyright (c) 2026 John Carter. All rights reserved.
"""One-shot fixture generator for the #179 e2e attachments suite.

The committed binary fixtures (``test.pdf`` / ``test.png`` /
``test.xlsx``) need to carry stable, model-recoverable content so the
e2e assertions stay deterministic across runs. Rebuild them with:

    uv run --with pillow --with openpyxl --with reportlab \
        python tests/e2e/fixtures/_generate.py

The generator never runs as part of the test suite — only when the
fixture content needs to change. Keep each output under 100 KB so the
repo doesn't bloat (issue #179 acceptance criterion).
"""

from __future__ import annotations

from pathlib import Path

HERE = Path(__file__).parent

# Strings the model must surface back to satisfy the e2e assertions.
PDF_PHRASE = "ELDERBERRY CONSORTIUM 4471"
XLSX_A1_VALUE = "PEPPERMINT"


def _write_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 24)
    c.drawString(72, 720, "Channel attachment e2e fixture (#179)")
    c.setFont("Helvetica", 16)
    c.drawString(72, 660, f"Phrase: {PDF_PHRASE}")
    c.drawString(72, 620, "This document is auto-generated; do not edit by hand.")
    c.showPage()
    c.save()


def _write_png(path: Path) -> None:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (128, 128), color=(248, 248, 248))
    draw = ImageDraw.Draw(img)
    # Solid red circle — distinct enough for a vision model to identify
    # the shape + colour without ambiguity.
    draw.ellipse((24, 24, 104, 104), fill=(220, 40, 40))
    img.save(path, format="PNG")


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = XLSX_A1_VALUE
    ws["B1"] = 42
    ws["C1"] = "control"
    wb.save(path)


def main() -> None:
    _write_pdf(HERE / "test.pdf")
    _write_png(HERE / "test.png")
    _write_xlsx(HERE / "test.xlsx")
    print("Fixtures regenerated:", sorted(p.name for p in HERE.glob("test.*")))


if __name__ == "__main__":
    main()
